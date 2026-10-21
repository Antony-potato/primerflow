"""Pruebas doradas: el motor reproduce las hojas del Excel de la tesis.

    python backend/test_motor.py        (sale con código 1 si algo falla)
"""
import re
import sys
from pathlib import Path

import openpyxl

import diseno
import motor
import pcr
import verificar

RAIZ = Path(__file__).resolve().parent.parent
EXCEL = RAIZ / "tonycheca15_resultados_esperados_por_manual1.xlsx"
FASTAS = sorted(p for p in RAIZ.glob("[0-9]_*.fasta"))  # los 9 del panel (sin el 6b opcional)
COND = {k: dict(v) for k, v in motor.PRESETS.items()}
# Erratas del Excel (µg/OD = PM / ε × 1000): 5447.6/154000 = 35.37 (dice 65.37); 6123.0/186800 = 32.78 (dice 30.78)
ERRATAS = {("Efaecium_ddl_cR", "ug_od"), ("Efaecalis_ddl_cR", "ug_od")}

if not EXCEL.exists():  # el Excel de la tesis no se publica en el repo
    sys.exit(f"Estas pruebas comparan contra {EXCEL.name}; ponlo en la raíz del proyecto para correrlas.")
_wb = openpyxl.load_workbook(EXCEL, read_only=True, data_only=True)


def hoja(nombre):
    return [r for r in _wb[nombre].iter_rows(min_row=5, values_only=True) if r[1] is not None]


def prefijo(archivo):
    """'3_S_aureus_nuc_15.fasta' → 'Saureus_nuc' (igual que el frontend)."""
    t = Path(archivo).stem.split("_")
    if re.fullmatch(r"\d+[a-z]?", t[0]):
        t = t[1:]
    return f"{t[0]}{t[1]}_{t[2]}" if len(t) >= 3 else Path(archivo).stem


def secuencias():
    out = []
    for f in FASTAS:
        out += [dict(s, especie=prefijo(f.name).split("_")[0]) for s in diseno.leer(f.read_text())]
    return out


def test_oligos_idt():
    # (campo del informe, columna del Excel, tolerancia; None = informativo)
    cols = [("tm_fabrica", 5, 0.11), ("tm_pcr", 6, None), ("pm", 7, 0.11), ("eps260", 8, 1),
            ("nmol_od", 9, 0.011), ("ug_od", 10, 0.011), ("dh", 11, 0.06), ("ds", 12, 0.06),
            ("horquilla_dg", 13, 1.5), ("horquilla_tm", 14, None), ("autodimero_dg_fabrica", 15, 0.011),
            ("autodimero_dg", 16, 0.011), ("dg_duplex_37", 17, 0.06)]
    filas = hoja("M3_IDT_oligos")
    inf = {r[1]: motor.informe(r[1], r[2], COND) for r in filas}
    for campo, col, tol in cols:
        err = {r[1]: abs((inf[r[1]][campo] or 0) - r[col]) for r in filas
               if r[col] is not None and (r[1], campo) not in ERRATAS}
        peor = max(err, key=err.get)
        print(f"   {campo:22s} error máx {err[peor]:6.2f} ({peor}){'' if tol else '  [calibrar]'}")
        assert tol is None or err[peor] <= tol, f"{campo}: {peor} se desvía {err[peor]:.2f}"
        assert len(filas) == 18


def test_verificador():
    filas = [r for r in _wb["M3_IDT_oligos"].iter_rows(min_row=4, values_only=True) if r[1] is not None]
    enc = [str(x) for x in filas[0] if x is not None]
    r = verificar.verificar_oligos(enc, [[str(x) for x in f[:len(enc)]] for f in filas[1:]], COND)
    malas = [(f[1]["tuyo"], h) for f in r["filas"] for h, c in zip(enc, f) if c["estado"] == "no"]
    assert sorted(malas) == [("Efaecalis_ddl_cR", "ug por OD260"), ("Efaecium_ddl_cR", "ug por OD260")], malas
    print(f"   {r['resumen']}")


def test_dimeros_idt():
    seq = {r[1]: r[2] for r in hoja("M3_IDT_oligos")}
    filas = hoja("M3_IDT_dimeros")
    for r in filas:
        a, b = seq[r[1]], seq[r[2]]
        assert abs(motor.dimero(a, b, COND["pcr"])["dg"] - r[4]) <= 0.011, r[:3]
        assert abs(motor.anclaje_3p(a, b, COND["pcr"]) - r[5]) <= 0.011, r[:3]
    assert len(filas) == 73


def test_conservacion_mega():
    esperado = {r[1]: r[4:12] for r in hoja("M2_archivos")}
    for f in FASTAS:
        a = diseno.alinear(diseno.leer(f.read_text()))
        res = a["resumen"]
        obtenido = (res["secuencias"], res["total"], res["con_hueco"], res["analizables"],
                    res["C"], res["V"], res["Pi"], res["S"])
        assert obtenido == tuple(esperado[f.name]), f"{f.name}: {obtenido} != {esperado[f.name]}"


def test_primers_en_alineamiento():
    alin = {f.name: diseno.alinear(diseno.leer(f.read_text())) for f in FASTAS}
    for r in hoja("M2_primers"):
        archivo, primer, motivo, ini, fin = r[0], r[3], r[6], r[7], r[8]
        for fila in alin[archivo]["filas"]:
            assert fila["seq"][ini - 1:fin] == motivo, f"{primer} no está en {fila['id']}"


def test_pcr_virtual():
    oligos = {r[1]: r[2] for r in hoja("M3_IDT_oligos")}
    pares = [{"nombre": n[:-3], "especie": n.split("_")[0], "F": s, "R": oligos[n[:-3] + "_cR"]}
             for n, s in oligos.items() if n.endswith("_cF")]
    res = {r["par"]: r for r in pcr.pcr_virtual(pares, secuencias(), cruzados=False)["resumen"]}
    for r in hoja("M4_PCR_virtual"):
        obtenido = res[r[0]]
        assert (obtenido["propia"], obtenido["otras"], obtenido["probadas_otras"]) == (r[2], r[3], r[4]), (r[0], obtenido)
        assert obtenido["tamanos"] == [int(x) for x in str(r[5]).split(",")], (r[0], obtenido["tamanos"])


def test_diseno_desde_fasta():
    f = next(p for p in FASTAS if "aureus" in p.name)
    a = diseno.alinear(diseno.leer(f.read_text()))
    d = diseno.disenar(a["plantilla"], a["conservacion"], "Saureus_nuc", COND)
    assert len(d["candidatos"]) >= 5, d["explicacion"]
    for p in d["candidatos"]:
        assert not p["rojos"] and 150 <= p["producto"]["tamano"] <= 600
        assert p["variables"] == 0  # 100 % conservado
        assert motor.rc(p["R"]["seq"]) == a["plantilla"][p["R"]["inicio"]:p["R"]["fin"]]
    print(f"   mejor par S. aureus: {d['candidatos'][0]['F']['seq']} / {d['candidatos'][0]['R']['seq']} "
          f"({d['candidatos'][0]['producto']['tamano']} pb, puntaje {d['candidatos'][0]['puntaje']})")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    fallos = 0
    for nombre, f in [(k, v) for k, v in globals().items() if k.startswith("test_")]:
        try:
            print(f"{nombre}")
            f()
            print("   OK")
        except AssertionError as e:
            fallos += 1
            print(f"   FALLA: {e}")
    sys.exit(1 if fallos else 0)

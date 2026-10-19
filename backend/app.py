"""API de PrimerFlow.  Correr:  uvicorn app:app --app-dir backend --port 8765"""
import io
import re
from pathlib import Path

from fastapi import FastAPI, Request
from fastapi.responses import JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

import diseno
import motor
import pcr
import verificar

app = FastAPI(title="PrimerFlow", version=motor.VERSION)


@app.exception_handler(ValueError)
def _error(_: Request, e: ValueError):
    return JSONResponse(status_code=422, content={"detail": str(e)})


class Condicion(BaseModel):
    oligo_uM: float = Field(gt=0, le=100)
    na_mM: float = Field(ge=0, le=2000)
    mg_mM: float = Field(ge=0, le=200)
    dntp_mM: float = Field(ge=0, le=50)


class Condiciones(BaseModel):
    fabrica: Condicion = Condicion(**motor.PRESETS["fabrica"])
    pcr: Condicion = Condicion(**motor.PRESETS["pcr"])
    multiplex: Condicion = Condicion(**motor.PRESETS["multiplex"])


class Oligo(BaseModel):
    nombre: str = Field(max_length=80)
    seq: str = Field(min_length=8, max_length=120)


class Plantilla(BaseModel):
    nombre: str = Field(max_length=80)
    seq: str = Field(min_length=20, max_length=200_000)


class Secuencia(BaseModel):
    id: str = Field(max_length=200)
    seq: str = Field(min_length=1, max_length=200_000)
    especie: str = Field("", max_length=80)


class AnalizarReq(BaseModel):
    oligos: list[Oligo] = Field(max_length=300)
    condiciones: Condiciones = Condiciones()


class DimerosReq(BaseModel):
    oligos: list[Oligo] = Field(min_length=1, max_length=60)
    condiciones: Condiciones = Condiciones()
    preset: str = Field("pcr", pattern="^(fabrica|pcr|multiplex)$")


class AlinearReq(BaseModel):
    secuencias: list[Secuencia] = Field(default=[], max_length=500)
    archivo: str | None = Field(None, max_length=20_000_000)  # .mul / .msf de MultAlin
    formato: str = Field("fasta", pattern="^(fasta|msf)$")
    ya_alineado: bool = False
    consenso_alto: float = Field(1.0, gt=0, le=1)
    consenso_bajo: float = Field(0.5, gt=0, le=1)


class DisenarReq(BaseModel):
    plantilla: str = Field(min_length=50, max_length=200_000)
    conservacion: list[float]
    prefijo: str = Field("diana", max_length=60)
    region: tuple[int, int] | None = None
    tamanos: list[tuple[int, int]] = [(150, 600)]
    excluir_tamanos: list[tuple[int, int]] = []
    oligos_tubo: list[Oligo] = Field(default=[], max_length=60)
    plantillas_tubo: list[Plantilla] = Field(default=[], max_length=30)  # consensos de las otras dianas del tubo
    condiciones: Condiciones = Condiciones()
    n: int = Field(20, ge=1, le=50)
    min_conservacion: int = Field(100, ge=0, le=100)


class Par(BaseModel):
    nombre: str = Field(max_length=80)
    especie: str = Field(max_length=80)
    tubo: str | None = Field(None, max_length=80)
    F: str = Field(min_length=8, max_length=120)
    R: str = Field(min_length=8, max_length=120)


class PcrReq(BaseModel):
    pares: list[Par] = Field(min_length=1, max_length=60)
    secuencias: list[Secuencia] = Field(min_length=1, max_length=2000)
    max_desajustes: int = Field(2, ge=0, le=6)
    tamano_max: int = Field(3000, ge=50, le=20_000)
    estricto_3p: bool = False
    cruzados: bool = True


class VerificarReq(BaseModel):
    encabezados: list[str] = Field(min_length=1, max_length=60)
    filas: list[list[str]] = Field(min_length=1, max_length=500)
    condiciones: Condiciones = Condiciones()


class Hoja(BaseModel):
    nombre: str = Field(min_length=1, max_length=60)
    encabezados: list[str] = Field(max_length=60)
    filas: list[list[str | float | int | None]] = Field(max_length=5000)


class ExcelReq(BaseModel):
    hojas: list[Hoja] = Field(min_length=1, max_length=10)


def _cond(c: Condiciones):
    return c.model_dump()


@app.get("/api/presets")
def presets():
    lim = lambda r: [None if abs(x) == motor.INF else x for x in r]  # JSON no admite inf: sin límite → null
    return {"motor": motor.VERSION, "presets": motor.PRESETS,
            "criterios": [{"regla": e, "campo": c, "verde": lim(v), "amarillo": lim(a)}
                          for e, c, v, a in motor.CRITERIOS_OLIGO + motor.CRITERIOS_PAR]}


@app.post("/api/analizar")
def analizar(req: AnalizarReq):
    cond = _cond(req.condiciones)
    return {"motor": motor.VERSION, "condiciones": cond,
            "oligos": [motor.informe(o.nombre, o.seq, cond) for o in req.oligos]}


@app.post("/api/dimeros")
def dimeros(req: DimerosReq):
    return motor.matriz_dimeros([o.model_dump() for o in req.oligos], _cond(req.condiciones)[req.preset])


@app.post("/api/alinear")
def alinear(req: AlinearReq):
    secs = diseno.leer(req.archivo, req.formato) if req.archivo else [s.model_dump() for s in req.secuencias]
    return diseno.alinear(secs, req.ya_alineado or bool(req.archivo), req.consenso_alto, req.consenso_bajo)


@app.post("/api/disenar")
def disenar(req: DisenarReq):
    return diseno.disenar(req.plantilla.upper(), req.conservacion, req.prefijo, _cond(req.condiciones),
                          req.region, req.tamanos, req.excluir_tamanos,
                          [o.model_dump() for o in req.oligos_tubo], req.n, req.min_conservacion,
                          [(p.nombre, p.seq.upper().replace("-", "")) for p in req.plantillas_tubo])


@app.post("/api/pcr-virtual")
def pcr_virtual(req: PcrReq):
    return pcr.pcr_virtual([p.model_dump() for p in req.pares], [s.model_dump() for s in req.secuencias],
                           req.max_desajustes, req.tamano_max, req.estricto_3p, req.cruzados)


@app.post("/api/verificar-oligos")
def verificar_oligos(req: VerificarReq):
    return verificar.verificar_oligos(req.encabezados, req.filas, _cond(req.condiciones))


@app.post("/api/excel")
def excel(req: ExcelReq):
    """Las tablas del analizador → un .xlsx con una hoja por tabla."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for h in req.hojas:
        ws = wb.create_sheet(re.sub(r"[\[\]:*?/\\]", " ", h.nombre)[:31])  # caracteres prohibidos en Excel
        ws.append(h.encabezados)
        for c in ws[1]:
            c.font = Font(bold=True)
        for f in h.filas:
            ws.append(f)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(42, max(len(str(c.value or "")) for c in col) + 2)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


dist = Path(__file__).resolve().parent.parent / "frontend" / "dist"
if dist.exists():
    app.mount("/", StaticFiles(directory=dist, html=True), name="frontend")
